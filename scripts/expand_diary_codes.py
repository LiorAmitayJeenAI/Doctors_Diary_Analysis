#!/usr/bin/env python3
"""Transform raw diary Excel exports into the anonymized-file format.

Finds the header row (even when a comment row sits above it), then:
  - keeps the canonical 24 appointment columns
  - expands short codes in מקור היומן, מעמד היומן, סטטוס זימון, סוג תור
  - reverses visual-order Hebrew in 15-תאור סיבה
  - normalizes header whitespace and ensures the comment row

Usage:
  python3 scripts/expand_diary_codes.py path/to/file.xlsx
  python3 scripts/expand_diary_codes.py path/to/file.xlsx -o path/to/out.xlsx
  python3 scripts/expand_diary_codes.py --in-place path/to/file.xlsx
  python3 scripts/expand_diary_codes.py --dry-run path/to/file.xlsx
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = REPO_ROOT / "data"

DEFAULT_FILES = [
    DATA_DIR / "אנונימי - תורים ליאוניד טיזנווה 01.06.2026.xlsx",
    DATA_DIR / "אנונימי - תורים מוריה גל מיתרים.xlsx",
]

HEADER_SCAN_ROWS = 10

COLUMN_MAPS: dict[str, dict[str, str]] = {
    "מקור היומן": {
        "נ": "נדחף",
        "ז": "לוז מקור",
        "ה": "הרחבה",
    },
    "מעמד היומן": {
        "מ": "ממלא מקום",
        "ס": "סגירה",
        "ע": "העדפה",
        "פ": "פנוי",
        "ר": "העדפה + ממלא מקום",
        "ש": "שבתון",
        "ה": "היעדרות",
        "א": "אודורו - להתעלם",
    },
    "סטטוס זימון": {
        "0": "0 - בוצע זימון / טרם בוצע ביקור",
        "1": "1 - בוצע ביקור",
        "3": "3 - חבר לא הגיע",
    },
    "סוג תור": {
        "0": "0 - רגיל",
        "1": "1 - נדחף",
    },
}

REQUIRED_HEADERS = ("מקור היומן", "מעמד היומן", "סטטוס זימון")

# Canonical 24 columns from the processed anonymized export.
# "שעה" appears twice: appointment time after תאריך, created time after יצירה תאריך.
KEEP_HEADERS = [
    ".ת.ז",
    "מתקן",
    "תפקיד",
    "תאריך",
    "שעה",
    "ק.זהות חבר",
    "ת.ז. חבר",
    "סוג תור",
    "סוג ביקור",
    "משך ביקור",
    "שעת ביקור בפועל(os)",
    "מקור היומן",
    "מעמד היומן",
    "קוד סיבה/סוג ביקור מועדף",
    "15-תאור סיבה",
    "סטטוס זימון",
    "segments-סהכ הוקצו לתור",
    "תאריך שיבוץ התור",
    "שעת שיבוץ התור",
    "קוד מתקן מוקד זימון",
    "יצירה תאריך",
    "שעה",
    "רשומה סטטוס",
    "תאריך שינוי אחרון",
]

HEBREW_COLUMN = "15-תאור סיבה"
HEBREW_CORRECT = {
    "עתודה",
    "סניף",
    "לקיחת דם",
    "מחלה",
    "תרגול החייאה",
}
HEBREW_REVERSED = {value[::-1] for value in HEBREW_CORRECT}

HEADER_COMMENTS = {
    "תאריך": "התור עצמו",
    "תאריך שיבוץ התור": "השעה שהחבר הזמין את התור - החתימה",
    "קוד מתקן מוקד זימון": "כל קוד אומר מוקד אחר - להתיחס כנתון",
    "יצירה תאריך": "חתימה שבה נקבע הסגמנט - הרופא פתח את התורים",
}


def normalize_header(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def code_key(value) -> str | None:
    """Turn a cell value into a lookup key, or None if the cell is empty."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text == "":
        return None
    if text.endswith(".0"):
        stem = text[:-2]
        if stem.isdigit() or (stem.startswith("-") and stem[1:].isdigit()):
            return stem
    return text


def has_hebrew(text: str) -> bool:
    return any("\u0590" <= ch <= "\u05FF" for ch in text)


def find_header_row(ws, required: tuple[str, ...] = REQUIRED_HEADERS) -> int | None:
    max_scan = min(HEADER_SCAN_ROWS, ws.max_row or 0)
    for row_idx in range(1, max_scan + 1):
        headers = [normalize_header(cell.value) for cell in ws[row_idx]]
        if all(name in headers for name in required):
            return row_idx
    return None


def header_names(ws, header_row: int) -> list[str]:
    max_col = ws.max_column or 0
    return [normalize_header(ws.cell(row=header_row, column=col).value) for col in range(1, max_col + 1)]


def normalize_header_row(ws, header_row: int) -> bool:
    changed = False
    for col_idx in range(1, (ws.max_column or 0) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        normalized = normalize_header(cell.value)
        if cell.value != normalized and not (cell.value is None and normalized == ""):
            cell.value = normalized or None
            changed = True
    return changed


def select_keep_columns(names: list[str]) -> list[int]:
    """Return 1-based indexes to keep, matching KEEP_HEADERS in order.

    Duplicate names such as שעה are bound left-to-right: first unused match wins.
    Headers that are missing from the sheet are skipped.
    """
    used: set[int] = set()
    keep: list[int] = []
    for wanted in KEEP_HEADERS:
        for idx, name in enumerate(names, start=1):
            if idx in used or name != wanted:
                continue
            keep.append(idx)
            used.add(idx)
            break
    return keep


def apply_mapping(ws, header_row: int, col_idx: int, mapping: dict[str, str]) -> tuple[dict, bool]:
    already_expanded = set(mapping.values())
    replacements: Counter[str] = Counter()
    empty = 0
    already = 0
    unknown: Counter[str] = Counter()
    changed = False

    for row_idx in range(header_row + 1, (ws.max_row or header_row) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        key = code_key(cell.value)
        if key is None:
            empty += 1
            continue
        if key in mapping:
            new_value = mapping[key]
            if cell.value != new_value:
                replacements[f"{key} → {new_value}"] += 1
                cell.value = new_value
                changed = True
            else:
                already += 1
        elif key in already_expanded:
            already += 1
        else:
            unknown[repr(cell.value)] += 1

    stats = {
        "replaced": dict(replacements),
        "replaced_total": sum(replacements.values()),
        "empty": empty,
        "already_expanded": already,
        "unknown": dict(unknown),
    }
    return stats, changed


def reverse_hebrew_column(ws, header_row: int, col_idx: int) -> tuple[dict, bool]:
    reversed_counts: Counter[str] = Counter()
    already = 0
    empty = 0
    unknown: Counter[str] = Counter()
    changed = False

    for row_idx in range(header_row + 1, (ws.max_row or header_row) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        key = code_key(cell.value)
        if key is None:
            empty += 1
            continue
        if key in HEBREW_CORRECT:
            already += 1
            continue
        if key in HEBREW_REVERSED or key[::-1] in HEBREW_CORRECT:
            new_value = key[::-1]
            reversed_counts[f"{key} → {new_value}"] += 1
            cell.value = new_value
            changed = True
            continue
        if has_hebrew(key):
            unknown[repr(cell.value)] += 1
        else:
            unknown[repr(cell.value)] += 1

    stats = {
        "replaced": dict(reversed_counts),
        "replaced_total": sum(reversed_counts.values()),
        "empty": empty,
        "already_expanded": already,
        "unknown": dict(unknown),
    }
    return stats, changed


def drop_columns_by_header(ws, header_row: int, keep_indexes: list[int]) -> list[str]:
    keep_set = set(keep_indexes)
    dropped: list[str] = []
    for col_idx in range(ws.max_column or 0, 0, -1):
        if col_idx in keep_set:
            continue
        dropped.append(normalize_header(ws.cell(row=header_row, column=col_idx).value) or f"col_{col_idx}")
        ws.delete_cols(col_idx)
    dropped.reverse()
    return dropped


def first_index_by_name(names: list[str], wanted: str) -> int | None:
    for idx, name in enumerate(names, start=1):
        if name == wanted:
            return idx
    return None


def ensure_comment_row(ws, header_row: int) -> tuple[int, dict, bool]:
    stats = {"inserted": False, "filled": {}}
    changed = False

    if header_row == 1:
        ws.insert_rows(1)
        header_row = 2
        stats["inserted"] = True
        changed = True

    comment_row = header_row - 1
    names = header_names(ws, header_row)
    filled: dict[str, str] = {}
    for header, comment in HEADER_COMMENTS.items():
        col_idx = first_index_by_name(names, header)
        if col_idx is None:
            continue
        cell = ws.cell(row=comment_row, column=col_idx)
        current = cell.value if cell.value is not None else ""
        if str(current).strip() == "":
            cell.value = comment
            filled[header] = comment
            changed = True
        elif normalize_header(current) != normalize_header(comment):
            # Leave an existing non-empty comment as-is.
            pass

    stats["filled"] = filled
    return header_row, stats, changed


def expand_workbook(path: Path, dry_run: bool = False, output_path: Path | None = None) -> dict:
    wb = openpyxl.load_workbook(path)
    dest = output_path or path
    maps = COLUMN_MAPS
    file_report = {
        "path": path,
        "output_path": dest,
        "sheets": [],
        "changed": False,
    }

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row = find_header_row(ws)
        if header_row is None:
            file_report["sheets"].append(
                {
                    "name": sheet_name,
                    "skipped": True,
                    "reason": "header row with diary columns not found",
                }
            )
            continue

        sheet_changed = False
        if normalize_header_row(ws, header_row):
            sheet_changed = True

        names = header_names(ws, header_row)
        keep_indexes = select_keep_columns(names)
        dropped = drop_columns_by_header(ws, header_row, keep_indexes)
        if dropped:
            sheet_changed = True

        names = header_names(ws, header_row)
        name_to_idx: dict[str, int] = {}
        for idx, name in enumerate(names, start=1):
            if name and name not in name_to_idx:
                name_to_idx[name] = idx
        # Second שעה: created-time column after יצירה תאריך is already kept by
        # select_keep_columns; mappings only need the first match by name.

        new_header_row, comment_stats, comment_changed = ensure_comment_row(ws, header_row)
        if comment_changed:
            sheet_changed = True
        header_row = new_header_row
        names = header_names(ws, header_row)
        name_to_idx = {}
        for idx, name in enumerate(names, start=1):
            if name and name not in name_to_idx:
                name_to_idx[name] = idx

        sheet_stats = {
            "name": sheet_name,
            "skipped": False,
            "header_row": header_row,
            "dropped_columns": dropped,
            "comment_row": comment_stats,
            "columns": {},
        }

        if HEBREW_COLUMN in name_to_idx:
            hebrew_stats, hebrew_changed = reverse_hebrew_column(
                ws, header_row, name_to_idx[HEBREW_COLUMN]
            )
            sheet_stats["columns"][HEBREW_COLUMN] = hebrew_stats
            if hebrew_changed:
                sheet_changed = True

        for col_name, mapping in maps.items():
            if col_name not in name_to_idx:
                continue
            if not mapping:
                continue
            col_stats, col_changed = apply_mapping(
                ws, header_row, name_to_idx[col_name], mapping
            )
            sheet_stats["columns"][col_name] = col_stats
            if col_changed:
                sheet_changed = True

        sheet_stats["changed"] = sheet_changed
        file_report["sheets"].append(sheet_stats)
        if sheet_changed:
            file_report["changed"] = True

    if not dry_run:
        dest.parent.mkdir(parents=True, exist_ok=True)
        wb.save(dest)
    wb.close()
    return file_report


def print_report(report: dict, dry_run: bool) -> None:
    path: Path = report["path"]
    dest: Path = report["output_path"]
    if dry_run:
        mode = "DRY-RUN"
    elif report["changed"]:
        mode = "SAVED"
    else:
        mode = "UNCHANGED"
    suffix = f" -> {dest.name}" if dest != path else ""
    print(f"\n=== {path.name} [{mode}]{suffix} ===")
    for sheet in report["sheets"]:
        if sheet.get("skipped"):
            print(f"  sheet {sheet['name']!r}: skipped ({sheet['reason']})")
            continue
        print(f"  sheet {sheet['name']!r}: header row {sheet['header_row']}")
        dropped = sheet.get("dropped_columns") or []
        if dropped:
            print(f"    dropped columns: {dropped}")
        comment = sheet.get("comment_row") or {}
        if comment.get("inserted"):
            print("    comment row: inserted")
        if comment.get("filled"):
            print(f"    comment row filled: {list(comment['filled'])}")
        for col_name, stats in sheet["columns"].items():
            print(f"    {col_name}:")
            print(f"      replaced: {stats['replaced_total']}")
            for label, count in stats["replaced"].items():
                print(f"        {count:5d}  {label}")
            print(f"      empty: {stats['empty']}")
            print(f"      already expanded: {stats['already_expanded']}")
            if stats["unknown"]:
                print("      unknown (left as-is):")
                for value, count in stats["unknown"].items():
                    print(f"        {count:5d}  {value}")


def default_output_path(source: Path) -> Path:
    return source.with_name(f"{source.stem}_processed{source.suffix}")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Transform diary Excel exports: expand codes, fix Hebrew, keep canonical columns."
    )
    parser.add_argument(
        "files",
        nargs="*",
        type=Path,
        help="Excel files to process. Defaults to the two current appointment exports.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print what would change without writing files.",
    )
    parser.add_argument(
        "--in-place",
        action="store_true",
        help="Overwrite the source file instead of writing *_processed.xlsx.",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Output path (only valid when a single input file is given).",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    files = args.files or list(DEFAULT_FILES)

    if args.in_place and args.output:
        print("Use either --in-place or -o, not both.", file=sys.stderr)
        return 2
    if args.output is not None and len(files) != 1:
        print("-o requires exactly one input file.", file=sys.stderr)
        return 2

    missing = [path for path in files if not path.exists()]
    if missing:
        for path in missing:
            print(f"File not found: {path}", file=sys.stderr)
        return 1

    for path in files:
        if args.output is not None:
            dest = args.output
        elif args.in_place:
            dest = path
        else:
            dest = default_output_path(path)
        report = expand_workbook(path, dry_run=args.dry_run, output_path=dest)
        print_report(report, dry_run=args.dry_run)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
